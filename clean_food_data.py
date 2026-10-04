# -*- coding: utf-8 -*-
"""
贵阳美食数据清洗脚本
输入: C:\\Users\\Administrator\\Desktop\\贵阳美食(1).xlsx
输出: C:\\Users\\Administrator\\Desktop\\贵阳美食_清洗后.xlsx

清洗规则:
1. 数字字段(评论数/平均消费/店铺评分/经度/纬度/各类id)从文本转为数字类型
2. 子评分 "口味:3.5 环境:3.6 服务:3.6" 拆分为 3 个独立列
3. 去除冗余列 "经纬度"(经度+纬度已单独存在)
4. 文本去首尾空格, 空字符串统一为 None
5. 空值保留原样(不删除行, 不填默认值, 避免丢失信息)
6. 店铺id 唯一, 不做去重
"""
import openpyxl
from openpyxl import Workbook
import re
import sys

SRC = r"C:\Users\Administrator\Desktop\贵阳美食(1).xlsx"
DST = r"C:\Users\Administrator\Desktop\贵阳美食_清洗后.xlsx"

# 原始列名 -> 列索引(0-based) 映射, 在读取表头后建立
# 需要转为数字的列
NUMERIC_COLS = {
    "评论数": int,
    "平均消费": float,
    "店铺评分": float,
    "经度": float,
    "纬度": float,
    "大类id": int,
    "中类id": int,
    "小类id": int,
    "城市代码": int,
}
# 需要删除的冗余列
DROP_COLS = {"经纬度"}
# 子评分拆分配置
SUB_SCORE_COL = "子评分"
SUB_SCORE_NEW = ["口味评分", "环境评分", "服务评分"]


def to_number(val, typ):
    """文本转数字, 失败返回 None"""
    if val is None:
        return None
    s = str(val).strip()
    if s == "" or s == "None":
        return None
    try:
        if typ is int:
            # 处理 "3.0" 这种带小数的整数
            f = float(s)
            return int(f) if f == int(f) else int(f)
        return float(s)
    except (ValueError, TypeError):
        return None


def clean_text(val):
    """文本清洗: 去首尾空格, 空字符串转 None"""
    if val is None:
        return None
    s = str(val).strip()
    return s if s != "" else None


def parse_sub_score(val):
    """拆分子评分 '口味:3.5 环境:3.6 服务:3.6' -> (3.5, 3.6, 3.6)"""
    if val is None:
        return (None, None, None)
    s = str(val).strip()
    result = [None, None, None]
    # 匹配 "口味:3.5" 等
    patterns = [r"口味[:：]\s*([\d.]+)", r"环境[:：]\s*([\d.]+)", r"服务[:：]\s*([\d.]+)"]
    for i, p in enumerate(patterns):
        m = re.search(p, s)
        if m:
            try:
                result[i] = float(m.group(1))
            except ValueError:
                pass
    return tuple(result)


def main():
    print(f"读取: {SRC}")
    wb_in = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws_in = wb_in["贵阳美食"]

    rows_iter = ws_in.iter_rows(values_only=True)
    header = list(next(rows_iter))
    print(f"原始表头({len(header)}列): {header}")

    # 建立列名 -> 索引
    col_idx = {}
    for i, h in enumerate(header):
        col_idx[h] = i

    # 构建新表头: 去掉冗余列, 在子评分后插入3个新列
    new_header = []
    for h in header:
        if h in DROP_COLS:
            continue
        new_header.append(h)
        if h == SUB_SCORE_COL:
            new_header.extend(SUB_SCORE_NEW)

    print(f"新表头({len(new_header)}列): {new_header}")

    # 创建输出工作簿(write_only 模式, 适合大数据量)
    wb_out = Workbook(write_only=True)
    ws_out = wb_out.create_sheet("贵阳美食")
    ws_out.append(new_header)

    total = 0
    numeric_idx = {col_idx[c]: t for c, t in NUMERIC_COLS.items() if c in col_idx}
    sub_score_i = col_idx.get(SUB_SCORE_COL)
    drop_idx = {col_idx[c] for c in DROP_COLS if c in col_idx}

    for row in rows_iter:
        total += 1
        if total % 10000 == 0:
            print(f"  已处理 {total} 行...")

        new_row = []
        for j, val in enumerate(row):
            if j in drop_idx:
                continue
            # 数字列转换
            if j in numeric_idx:
                new_row.append(to_number(val, numeric_idx[j]))
                continue
            # 文本列清洗
            cleaned = clean_text(val)
            new_row.append(cleaned)
            # 子评分拆分
            if j == sub_score_i:
                kou, env, serv = parse_sub_score(val)
                new_row.extend([kou, env, serv])

        ws_out.append(new_row)

    wb_in.close()
    print(f"写入: {DST}")
    wb_out.save(DST)
    print(f"完成! 共处理 {total} 行, 输出 {len(new_header)} 列")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"错误: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)
